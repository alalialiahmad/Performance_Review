"""
exporter.py — Excel export with color-coded KPI sheets and PNG chart export.

Creates multi-sheet Excel workbooks with conditional color formatting
matching the KPI thresholds (green/yellow/red).
"""

import logging
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from analytics import get_kpi_color, KPI_TARGETS, REASON_DISPLAY_NAMES

logger = logging.getLogger(__name__)

# Color definitions matching KPI status
FILL_GREEN = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
FILL_YELLOW = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
FILL_RED = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
FILL_HEADER = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

FONT_HEADER = Font(color="FFFFFF", bold=True, size=11)
FONT_TITLE = Font(bold=True, size=14)
FONT_BODY = Font(size=11)

ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")

THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

COLOR_MAP = {"green": FILL_GREEN, "yellow": FILL_YELLOW, "red": FILL_RED}

# KPI columns to color-code in the Per-Camera sheet
KPI_COLUMNS_TO_COLOR = list(KPI_TARGETS.keys())


def get_fill_for_kpi(kpi_name: str, value: float) -> PatternFill | None:
    """Map a KPI value to its color fill based on thresholds."""
    color = get_kpi_color(kpi_name, value)
    return COLOR_MAP.get(color)


def export_to_excel(
    file_path: str,
    fleet_kpis: dict,
    camera_kpis_df: pd.DataFrame,
    manual_reasons_df: pd.DataFrame,
    outlier_report_df: pd.DataFrame,
) -> str:
    """
    Create a 4-sheet Excel workbook with color-coded KPI cells.

    Sheets:
        1. Fleet Summary — fleet-wide KPI values with colors
        2. Per-Camera KPIs — one row per camera, KPI cells colored
        3. Manual Reasons Detail — breakdown of 7 review reasons
        4. Outlier Report — flagged cameras with reasons

    Args:
        file_path: Output path for the .xlsx file
        fleet_kpis: Dict from analytics.compute_fleet_kpis()
        camera_kpis_df: DataFrame from analytics.compute_per_camera_kpis()
        manual_reasons_df: DataFrame from analytics.compute_manual_review_breakdown()
        outlier_report_df: DataFrame from outliers.generate_outlier_report()

    Returns:
        The file path of the saved workbook.
    """
    wb = Workbook()

    _write_fleet_summary_sheet(wb, fleet_kpis)
    _write_camera_kpis_sheet(wb, camera_kpis_df)
    _write_reasons_sheet(wb, manual_reasons_df)
    _write_outlier_sheet(wb, outlier_report_df)

    wb.save(file_path)
    logger.info("Exported Excel report to %s", file_path)
    return str(file_path)


def _write_fleet_summary_sheet(wb: Workbook, summary: dict) -> None:
    """Write fleet-wide KPI summary as a vertical label-value table."""
    ws = wb.active
    ws.title = "Fleet Summary"

    # Title
    ws.merge_cells("A1:C1")
    ws["A1"] = "Fleet-Wide KPI Summary"
    ws["A1"].font = FONT_TITLE

    # Headers
    headers = ["KPI", "Value", "Status"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col, value=header)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER

    # KPI rows
    kpi_display = [
        ("accuracy_pct", "Camera Accuracy %"),
        ("manual_review_rate_pct", "Manual Review Rate %"),
        ("linking_rate_pct", "Camera Linking Rate %"),
        ("pending_rate_pct", "Pending Rate %"),
        ("ocr_review_rate_pct", "OCR Review Rate %"),
        ("manual_link_rate_pct", "Manual Link Rate %"),
        ("archive_rate_pct", "Archive Rate %"),
        ("archive_by_reviewer_rate_pct", "Archive by Reviewer %"),
        ("archive_by_ops_rate_pct", "Archive by Ops %"),
        ("archive_by_auto_rate_pct", "Archive by Auto %"),
    ]

    row = 4
    for kpi_key, display_name in kpi_display:
        value = summary.get(kpi_key, 0.0)

        ws.cell(row=row, column=1, value=display_name).border = THIN_BORDER
        val_cell = ws.cell(row=row, column=2, value=f"{value:.1f}%")
        val_cell.alignment = ALIGN_CENTER
        val_cell.border = THIN_BORDER

        # Color-code and status
        fill = get_fill_for_kpi(kpi_key, value)
        status = summary.get(f"{kpi_key}_status", "")
        status_cell = ws.cell(row=row, column=3, value=status.upper() if status else "")
        status_cell.alignment = ALIGN_CENTER
        status_cell.border = THIN_BORDER

        if fill:
            val_cell.fill = fill
            status_cell.fill = fill

        row += 1

    # Counts section
    row += 1
    ws.cell(row=row, column=1, value="Total Detected").border = THIN_BORDER
    ws.cell(row=row, column=2, value=int(summary.get("total_detected", 0))).border = THIN_BORDER
    row += 1
    ws.cell(row=row, column=1, value="Total Transactions").border = THIN_BORDER
    ws.cell(row=row, column=2, value=int(summary.get("total_transactions", 0))).border = THIN_BORDER
    row += 1
    ws.cell(row=row, column=1, value="Total Archived").border = THIN_BORDER
    ws.cell(row=row, column=2, value=int(summary.get("total_archived", 0))).border = THIN_BORDER

    _auto_fit_columns(ws)


def _write_camera_kpis_sheet(wb: Workbook, camera_df: pd.DataFrame) -> None:
    """Write per-camera KPI table with color-coded cells."""
    ws = wb.create_sheet("Per-Camera KPIs")

    if camera_df.empty:
        ws["A1"] = "No camera data available"
        return

    # Select display columns (exclude internal color/status columns)
    display_cols = [
        "camera_name", "lot_name", "camera_direction", "total_detected",
        "accuracy_pct", "manual_review_rate_pct", "linking_rate_pct",
        "pending_rate_pct", "ocr_review_rate_pct", "manual_link_rate_pct",
        "archive_rate_pct", "archive_by_reviewer_rate_pct", "archive_by_ops_rate_pct",
    ]
    display_cols = [c for c in display_cols if c in camera_df.columns]

    display_headers = {
        "camera_name": "Camera Name",
        "lot_name": "Lot Name",
        "camera_direction": "Direction",
        "total_detected": "Total Detected",
        "accuracy_pct": "Accuracy %",
        "manual_review_rate_pct": "Manual Review %",
        "linking_rate_pct": "Linking %",
        "pending_rate_pct": "Pending %",
        "ocr_review_rate_pct": "OCR Review %",
        "manual_link_rate_pct": "Manual Link %",
        "archive_rate_pct": "Archive %",
        "archive_by_reviewer_rate_pct": "Archive by Reviewer %",
        "archive_by_ops_rate_pct": "Archive by Ops %",
    }

    # Write headers
    for col_idx, col_name in enumerate(display_cols, 1):
        cell = ws.cell(row=1, column=col_idx, value=display_headers.get(col_name, col_name))
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER

    # Write data rows
    for row_idx, (_, row) in enumerate(camera_df.iterrows(), 2):
        for col_idx, col_name in enumerate(display_cols, 1):
            value = row.get(col_name, "")

            if isinstance(value, float) and col_name != "total_detected":
                cell = ws.cell(row=row_idx, column=col_idx, value=f"{value:.1f}%")
            elif col_name == "total_detected":
                cell = ws.cell(row=row_idx, column=col_idx, value=int(value) if pd.notna(value) else 0)
            else:
                cell = ws.cell(row=row_idx, column=col_idx, value=str(value))

            cell.border = THIN_BORDER
            cell.alignment = ALIGN_CENTER if col_idx > 3 else ALIGN_LEFT

            # Color-code KPI cells
            if col_name in KPI_COLUMNS_TO_COLOR:
                fill = get_fill_for_kpi(col_name, float(value) if pd.notna(value) else 0)
                if fill:
                    cell.fill = fill

    _auto_fit_columns(ws)


def _write_reasons_sheet(wb: Workbook, reasons_df: pd.DataFrame) -> None:
    """Write manual review reasons breakdown table."""
    ws = wb.create_sheet("Manual Reasons Detail")

    if reasons_df.empty:
        ws["A1"] = "No manual review data available"
        return

    headers = ["Reason", "Count", "% of Manual Review", "% of Total Detected"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER

    for row_idx, (_, row) in enumerate(reasons_df.iterrows(), 2):
        ws.cell(row=row_idx, column=1, value=row.get("display_name", "")).border = THIN_BORDER
        ws.cell(row=row_idx, column=2, value=int(row.get("count", 0))).border = THIN_BORDER

        pct_review = row.get("pct_of_review", 0)
        cell = ws.cell(row=row_idx, column=3, value=f"{pct_review:.1f}%")
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER
        if pct_review > 50:
            cell.fill = FILL_RED

        pct_detected = row.get("pct_of_detected", 0)
        cell = ws.cell(row=row_idx, column=4, value=f"{pct_detected:.1f}%")
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER

    _auto_fit_columns(ws)


def _write_outlier_sheet(wb: Workbook, outlier_df: pd.DataFrame) -> None:
    """Write outlier report table with color-coded rows."""
    ws = wb.create_sheet("Outlier Report")

    if outlier_df.empty:
        ws["A1"] = "No outliers detected"
        return

    headers = ["Camera", "Lot", "Metric", "Value", "Fleet Mean", "Z-Score", "Flag Type", "Reason"]
    col_keys = ["camera_name", "lot_name", "metric", "value", "fleet_mean", "z_score", "flag_type", "reason"]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER

    # Row color by flag type
    flag_colors = {
        "Threshold Breach": FILL_RED,
        "Z-Score": FILL_YELLOW,
        "IQR": FILL_YELLOW,
        "Dominant Reason": PatternFill(start_color="E8DAEF", end_color="E8DAEF", fill_type="solid"),
    }

    for row_idx, (_, row) in enumerate(outlier_df.iterrows(), 2):
        flag_type = row.get("flag_type", "")
        row_fill = flag_colors.get(flag_type)

        for col_idx, key in enumerate(col_keys, 1):
            value = row.get(key, "")
            if isinstance(value, float):
                cell = ws.cell(row=row_idx, column=col_idx, value=f"{value:.1f}" if value != 0 else "-")
            else:
                cell = ws.cell(row=row_idx, column=col_idx, value=str(value))

            cell.border = THIN_BORDER
            if row_fill:
                cell.fill = row_fill

    _auto_fit_columns(ws)


def _auto_fit_columns(ws) -> None:
    """Auto-fit column widths based on content."""
    for col_cells in ws.columns:
        max_length = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_length + 3, 50)


def export_dashboard_png(figure, file_path: str) -> str:
    """
    Save a matplotlib Figure as a PNG image.

    Args:
        figure: matplotlib.figure.Figure object
        file_path: Output path for the PNG

    Returns:
        The file path of the saved image.
    """
    figure.savefig(file_path, dpi=150, bbox_inches="tight", facecolor=figure.get_facecolor())
    logger.info("Exported dashboard PNG to %s", file_path)
    return str(file_path)
